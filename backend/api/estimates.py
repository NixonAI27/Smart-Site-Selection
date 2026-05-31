import io
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from api.auth import get_current_user
from models.database import get_db
from models.estimate import Estimate
from models.line_item import LineItem
from models.project import Project

router = APIRouter(prefix="/api/estimates", tags=["estimates"])


class LineItemCreate(BaseModel):
    category: str = "other"
    description: str
    quantity: float = 1.0
    unit: str = "ea"
    unit_cost: float = 0.0
    ai_generated: bool = False


class EstimateCreate(BaseModel):
    project_id: int
    markup_pct: float = 15.0
    notes: str | None = None
    line_items: list[LineItemCreate] = []


class EstimateUpdate(BaseModel):
    status: str | None = None
    markup_pct: float | None = None
    notes: str | None = None


def _calc_totals(estimate: Estimate) -> None:
    labor = sum(i.total_cost for i in estimate.line_items if i.category == "labor")
    materials = sum(i.total_cost for i in estimate.line_items if i.category != "labor")
    estimate.total_labor = labor
    estimate.total_materials = materials
    subtotal = labor + materials
    estimate.grand_total = subtotal * (1 + estimate.markup_pct / 100)


def _estimate_dict(e: Estimate) -> dict:
    return {
        "id": e.id,
        "project_id": e.project_id,
        "version": e.version,
        "status": e.status,
        "total_labor": e.total_labor,
        "total_materials": e.total_materials,
        "markup_pct": e.markup_pct,
        "grand_total": e.grand_total,
        "notes": e.notes,
        "created_at": e.created_at.isoformat() if e.created_at else None,
        "line_items": [_item_dict(i) for i in e.line_items],
    }


def _item_dict(i: LineItem) -> dict:
    return {
        "id": i.id,
        "category": i.category,
        "description": i.description,
        "quantity": i.quantity,
        "unit": i.unit,
        "unit_cost": i.unit_cost,
        "total_cost": i.total_cost,
        "ai_generated": i.ai_generated,
    }


@router.get("/project/{project_id}")
def list_estimates(project_id: int, db: Session = Depends(get_db), _: str = Depends(get_current_user)):
    return [_estimate_dict(e) for e in db.query(Estimate).filter(Estimate.project_id == project_id).all()]


@router.post("", status_code=201)
def create_estimate(body: EstimateCreate, db: Session = Depends(get_db), _: str = Depends(get_current_user)):
    proj = db.query(Project).filter(Project.id == body.project_id).first()
    if not proj:
        raise HTTPException(404, "Project not found")
    # auto-increment version
    existing = db.query(Estimate).filter(Estimate.project_id == body.project_id).count()
    est = Estimate(project_id=body.project_id, version=existing + 1, markup_pct=body.markup_pct, notes=body.notes)
    db.add(est)
    db.flush()
    for item_data in body.line_items:
        item = LineItem(estimate_id=est.id, **item_data.model_dump())
        item.total_cost = item.quantity * item.unit_cost
        db.add(item)
    db.flush()
    db.refresh(est)
    _calc_totals(est)
    db.commit()
    db.refresh(est)
    return _estimate_dict(est)


@router.get("/{estimate_id}")
def get_estimate(estimate_id: int, db: Session = Depends(get_db), _: str = Depends(get_current_user)):
    e = db.query(Estimate).filter(Estimate.id == estimate_id).first()
    if not e:
        raise HTTPException(404, "Estimate not found")
    return _estimate_dict(e)


@router.patch("/{estimate_id}")
def update_estimate(estimate_id: int, body: EstimateUpdate, db: Session = Depends(get_db), _: str = Depends(get_current_user)):
    e = db.query(Estimate).filter(Estimate.id == estimate_id).first()
    if not e:
        raise HTTPException(404, "Estimate not found")
    for k, v in body.model_dump(exclude_none=True).items():
        setattr(e, k, v)
    _calc_totals(e)
    db.commit()
    db.refresh(e)
    return _estimate_dict(e)


@router.delete("/{estimate_id}", status_code=204)
def delete_estimate(estimate_id: int, db: Session = Depends(get_db), _: str = Depends(get_current_user)):
    e = db.query(Estimate).filter(Estimate.id == estimate_id).first()
    if not e:
        raise HTTPException(404, "Estimate not found")
    db.delete(e)
    db.commit()


# --- Line item CRUD ---

class LineItemUpdate(BaseModel):
    category: str | None = None
    description: str | None = None
    quantity: float | None = None
    unit: str | None = None
    unit_cost: float | None = None


@router.post("/{estimate_id}/items", status_code=201)
def add_line_item(estimate_id: int, body: LineItemCreate, db: Session = Depends(get_db), _: str = Depends(get_current_user)):
    e = db.query(Estimate).filter(Estimate.id == estimate_id).first()
    if not e:
        raise HTTPException(404, "Estimate not found")
    item = LineItem(estimate_id=estimate_id, **body.model_dump())
    item.total_cost = item.quantity * item.unit_cost
    db.add(item)
    db.flush()
    db.refresh(e)
    _calc_totals(e)
    db.commit()
    db.refresh(item)
    return _item_dict(item)


@router.patch("/{estimate_id}/items/{item_id}")
def update_line_item(estimate_id: int, item_id: int, body: LineItemUpdate, db: Session = Depends(get_db), _: str = Depends(get_current_user)):
    item = db.query(LineItem).filter(LineItem.id == item_id, LineItem.estimate_id == estimate_id).first()
    if not item:
        raise HTTPException(404, "Line item not found")
    for k, v in body.model_dump(exclude_none=True).items():
        setattr(item, k, v)
    item.total_cost = item.quantity * item.unit_cost
    e = db.query(Estimate).filter(Estimate.id == estimate_id).first()
    db.flush()
    db.refresh(e)
    _calc_totals(e)
    db.commit()
    db.refresh(item)
    return _item_dict(item)


@router.delete("/{estimate_id}/items/{item_id}", status_code=204)
def delete_line_item(estimate_id: int, item_id: int, db: Session = Depends(get_db), _: str = Depends(get_current_user)):
    item = db.query(LineItem).filter(LineItem.id == item_id, LineItem.estimate_id == estimate_id).first()
    if not item:
        raise HTTPException(404, "Line item not found")
    db.delete(item)
    e = db.query(Estimate).filter(Estimate.id == estimate_id).first()
    db.flush()
    db.refresh(e)
    _calc_totals(e)
    db.commit()


# --- Export ---

@router.get("/{estimate_id}/export")
def export_estimate(estimate_id: int, format: str = "pdf", db: Session = Depends(get_db), _: str = Depends(get_current_user)):
    e = db.query(Estimate).filter(Estimate.id == estimate_id).first()
    if not e:
        raise HTTPException(404, "Estimate not found")
    proj = db.query(Project).filter(Project.id == e.project_id).first()

    if format == "xlsx":
        return _export_xlsx(e, proj)
    return _export_pdf(e, proj)


def _export_pdf(estimate: Estimate, project: Project) -> StreamingResponse:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(f"Finish Carpentry Estimate", styles["Title"]))
    if project:
        story.append(Paragraph(f"Project: {project.name} | Client: {project.client or '-'}", styles["Normal"]))
        story.append(Paragraph(f"Address: {project.address or '-'}", styles["Normal"]))
    story.append(Paragraph(f"Status: {estimate.status.upper()} | Version: {estimate.version}", styles["Normal"]))
    story.append(Spacer(1, 12))

    data = [["Category", "Description", "Qty", "Unit", "Unit Cost", "Total"]]
    for item in estimate.line_items:
        data.append([
            item.category, item.description,
            f"{item.quantity:.2f}", item.unit,
            f"${item.unit_cost:,.2f}", f"${item.total_cost:,.2f}",
        ])
    data.append(["", "", "", "", "Labor", f"${estimate.total_labor:,.2f}"])
    data.append(["", "", "", "", "Materials", f"${estimate.total_materials:,.2f}"])
    data.append(["", "", "", "", f"Markup ({estimate.markup_pct:.0f}%)", ""])
    data.append(["", "", "", "", "GRAND TOTAL", f"${estimate.grand_total:,.2f}"])

    t = Table(data, colWidths=[70, 200, 40, 40, 70, 80])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2C4A2A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -5), [colors.white, colors.HexColor("#F2EDE3")]),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    story.append(t)
    doc.build(story)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename=estimate_{estimate.id}.pdf"})


def _export_xlsx(estimate: Estimate, project: Project) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Estimate"

    header_fill = PatternFill("solid", fgColor="2C4A2A")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append(["Finish Carpentry Estimate"])
    if project:
        ws.append([f"Project: {project.name}", f"Client: {project.client or '-'}", f"Address: {project.address or '-'}"])
    ws.append([])

    headers = ["Category", "Description", "Quantity", "Unit", "Unit Cost", "Total Cost"]
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = header_font
        cell.fill = header_fill

    for item in estimate.line_items:
        ws.append([item.category, item.description, item.quantity, item.unit, item.unit_cost, item.total_cost])

    ws.append([])
    ws.append(["", "", "", "", "Total Labor", estimate.total_labor])
    ws.append(["", "", "", "", "Total Materials", estimate.total_materials])
    ws.append(["", "", "", "", f"Markup {estimate.markup_pct:.0f}%", ""])
    ws.append(["", "", "", "", "GRAND TOTAL", estimate.grand_total])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf,
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=estimate_{estimate.id}.xlsx"})
