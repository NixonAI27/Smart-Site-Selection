import base64
import io
import json
from pathlib import Path
from typing import Any


def parse_document(file_path: str, filename: str) -> dict[str, Any]:
    """Parse a document and return text content + optional vision images."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _parse_pdf(file_path)
    elif ext in (".docx",):
        return _parse_docx(file_path)
    elif ext in (".xlsx", ".xls"):
        return _parse_excel(file_path)
    elif ext == ".dxf":
        return _parse_dxf(file_path)
    elif ext in (".jpg", ".jpeg", ".png"):
        return _parse_image(file_path, ext)
    else:
        return {"text": "", "images": [], "error": f"Unsupported file type: {ext}"}


def _parse_pdf(file_path: str) -> dict:
    import pdfplumber
    import fitz  # pymupdf

    text_parts = []
    images = []

    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    page_text += " | ".join(str(c) for c in row if c) + "\n"
            text_parts.append(page_text)

    doc = fitz.open(file_path)
    for page_num in range(min(len(doc), 10)):  # cap at 10 pages for vision
        page = doc[page_num]
        mat = fitz.Matrix(1.5, 1.5)
        pix = page.get_pixmap(matrix=mat)
        img_bytes = pix.tobytes("png")
        images.append({
            "type": "image",
            "source": {
                "type": "base64",
                "media_type": "image/png",
                "data": base64.standard_b64encode(img_bytes).decode(),
            },
        })
    doc.close()

    return {"text": "\n\n".join(text_parts), "images": images}


def _parse_docx(file_path: str) -> dict:
    from docx import Document

    doc = Document(file_path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return {"text": "\n".join(parts), "images": []}


def _parse_excel(file_path: str) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return {"text": "\n".join(parts), "images": []}


def _parse_dxf(file_path: str) -> dict:
    import ezdxf

    doc = ezdxf.readfile(file_path)
    msp = doc.modelspace()
    entities = []
    for e in msp:
        if e.dxftype() == "TEXT":
            entities.append({"type": "TEXT", "value": e.dxf.text, "pos": list(e.dxf.insert[:2])})
        elif e.dxftype() == "MTEXT":
            entities.append({"type": "MTEXT", "value": e.text})
        elif e.dxftype() == "LINE":
            start = list(e.dxf.start[:2])
            end = list(e.dxf.end[:2])
            length = ((end[0]-start[0])**2 + (end[1]-start[1])**2) ** 0.5
            entities.append({"type": "LINE", "length": round(length, 2)})
        elif e.dxftype() == "DIMENSION":
            entities.append({"type": "DIMENSION", "text": getattr(e.dxf, "text", "")})
    return {"text": json.dumps(entities, indent=2), "images": []}


def _parse_image(file_path: str, ext: str) -> dict:
    from PIL import Image

    media_type = "image/jpeg" if ext in (".jpg", ".jpeg") else "image/png"
    with Image.open(file_path) as img:
        # resize large images to keep under API limits
        max_dim = 1568
        if max(img.width, img.height) > max_dim:
            ratio = max_dim / max(img.width, img.height)
            img = img.resize((int(img.width * ratio), int(img.height * ratio)), Image.LANCZOS)
        buf = io.BytesIO()
        save_fmt = "JPEG" if ext in (".jpg", ".jpeg") else "PNG"
        img.save(buf, format=save_fmt)
        img_b64 = base64.standard_b64encode(buf.getvalue()).decode()

    image_block = {
        "type": "image",
        "source": {"type": "base64", "media_type": media_type, "data": img_b64},
    }
    return {"text": "", "images": [image_block]}
